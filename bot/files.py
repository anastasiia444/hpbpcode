"""Резюме, надіслані в Telegram: зберігаємо файл і готуємо його для Claude."""

import base64
import re
from datetime import datetime
from pathlib import Path

from docx import Document

MAX_FILE_BYTES = 20 * 1024 * 1024  # Telegram Bot API не віддає більші файли
SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def safe_filename(original: str) -> str:
    stem, suffix = Path(original).stem, Path(original).suffix.lower()
    stem = re.sub(r"[^\w\-]+", "_", stem, flags=re.UNICODE).strip("_")[:60] or "resume"
    return f"{datetime.now():%Y%m%d-%H%M%S}-{stem}{suffix}"


def docx_text(path: Path) -> str:
    document = Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def resume_content_blocks(path: Path) -> list[dict]:
    """Блоки повідомлення для Claude з вмістом файлу."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        return [
            {
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": data},
                "title": path.name,
            }
        ]
    if suffix == ".docx":
        text = docx_text(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        raise ValueError("У файлі не знайшлося тексту")
    return [{"type": "text", "text": f'<resume file="{path.name}">\n{text}\n</resume>'}]
